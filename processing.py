import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from utils import *  # Assuming all processing functions are defined in utils.py
from database import fetch_client_load_profile
import pandas as pd  # Import pandas for DataFrame manipulation

# Define the relative paths for images and PowerPoint template
image_folder = 'image'
template_path = 'solx.pptx'


def add_table_to_slide(slide, enery_summary, rows, cols, x, y, width, height):
  """
  Adds a table to a PowerPoint slide.
  """
  table = slide.shapes.add_table(rows, cols, x, y, width, height).table

  # Adjust column width for the first column
  table.columns[0].width = Inches(2.2)

  # Define the row header color (using RGB)
  row_header_color = RGBColor(249, 163, 28)

  # Rename the columns for readability
  new_column_names = {
      'supply_period': 'Supply Period',
      'number of hours': 'Total Hours',
      'kwh': 'Consumption(kWh)',
      'kw': 'Demand(kW)',
      'load factor': 'Load Factor'
  }

  # Ensure column names are in string format for consistency
  enery_summary.columns = enery_summary.columns.astype(str)

  # Add header with orange background and white font
  for col_idx, col_name in enumerate(enery_summary.columns):
    new_name = new_column_names.get(col_name.lower(), col_name)
    cell = table.cell(0, col_idx)
    cell.text = str(new_name)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(249, 163, 28)  # Orange
    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)  # White
    cell.text_frame.paragraphs[0].font.size = Pt(12)
    cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

  # Fill table with data and adjust font size for each cell
  for row_idx, row_data in enumerate(enery_summary.values):
    for col_idx, cell_value in enumerate(row_data):
      cell = table.cell(row_idx + 1, col_idx)  # +1 to skip the header row

      # Convert numeric values to string format
      if isinstance(cell_value, (int, float)):
        cell_value = f"{cell_value:,.2f}"  # Format numeric values to 2 decimal places
      else:
        cell_value = str(cell_value)  # Convert non-numeric values to string

      cell.text = str(cell_value)
      cell.text_frame.paragraphs[0].font.size = Pt(10)  # Font size for content
      cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Bold the rows for 'Total', 'Average', 'Max', 'Min'
    if row_data[0] in ['Total', 'Average', 'Max', 'Min']:
      for col_idx in range(cols):
        cell = table.cell(row_idx + 1, col_idx)
        cell.text_frame.paragraphs[0].font.bold = True

  return table


def process_data_and_generate_reports(client_name: str):
  """
  Fetches and processes client data, generating various energy-related reports.
  """
  try:
    # Fetch client data from the database
    client_data = fetch_client_load_profile(client_name)

    if client_data.empty:
      raise ValueError(f"No data found for client: {client_name}")

    # Ensure all necessary columns are in the correct format
    # client_data['kwh'] = pd.to_numeric(client_data['kwh'], errors='coerce')
    # client_data['kw'] = pd.to_numeric(client_data['kw'], errors='coerce')

    # Process data and generate summaries and plots
    processed_data = process_energy_data(client_data)
    energy_summary = generate_energy_summary(processed_data)

    # Check if 'load factor' is in the summary
    if 'load factor' not in energy_summary.columns:
      print(f"Warning: 'load factor' missing from energy summary. Adding default NaN values.")
      energy_summary['load factor'] = pd.NA  # Adding missing load factor as NaN

    print("Energy Summary Columns:", energy_summary.columns.tolist())

    # Create image folder if it doesn't exist
    image_folder = "image"
    os.makedirs(image_folder, exist_ok=True)

    # Generate images with debugging
    try:
      print("Generating energy consumption plot...")
      energy_consumption_plot(processed_data)
      print("Energy consumption plot saved.")
    except Exception as e:
      print(f"Error generating energy consumption plot: {e}")

    try:
      print("Generating energy behavior plot...")
      energy_behavior_plot(processed_data)
      print("Energy behavior plot saved.")
    except Exception as e:
      print(f"Error generating energy behavior plot: {e}")

    # Hourly load curve
    try:
      hourly_load_curve_by_day = client_data.pivot_table(
          index="hour",
          columns="weekday",
          values="kw",
          aggfunc={"kw": "mean"}
      )
      plot_hourly_by_day_load_curve(hourly_load_curve_by_day, column_name='max', unit='kW')
      print("Hourly load curve plot saved.")
    except Exception as e:
      print(f"Error generating hourly load curve plot: {e}")

    try:
      hourly_load_table(hourly_load_curve_by_day)
      print("Hourly heatmap saved.")
    except Exception as e:
      print(f"Erorr generating hourly heatmap: {e}")

    # Demand and consumption
    try:
      plot_demand_and_consumption_WESM(processed_data)
      print("Demand and consumption plot saved.")
    except Exception as e:
      print(f"Error generating demand and consumption plot: {e}")

    # Daily report
    try:
      generate_daily_report(
          client_data=client_data,
          values_column='kw',
          wesm_column='wesm'
      )
      print("Daily report saved.")
    except Exception as e:
      print(f"Error generating daily report: {e}")

    # Return the processed data and images for PowerPoint creation
    image_paths = {
        'energy_consumption': os.path.join(image_folder, 'energy_consumption.png'),
        'energy_behavior': os.path.join(image_folder, 'energy_behavior.png'),
        'hourly_load_curve': os.path.join(image_folder, 'hourly_load_curve.png'),
        'peak_demand_plot': os.path.join(image_folder, 'peak_demand_plot.png'),
        'demand_consumption_wesm': os.path.join(image_folder, 'demand_consumption_wesm.png'),
        'daily_kw_report': os.path.join(image_folder, 'daily_kw_report.png'),
    }

    return energy_summary, image_paths

  except Exception as e:
    print(f"Error in process_data_and_generate_reports: {e}")
    return None, None


def generate_ppt(client_name: str, image_folder: str, template_path: str):
  """
  Main function to generate PowerPoint report for the given client.
  """
  try:
    energy_summary, image_paths = process_data_and_generate_reports(client_name)

    # Load PowerPoint template
    ppt = Presentation(template_path)

    # Create slides and add content
    slide1 = ppt.slides.add_slide(ppt.slide_layouts[0])  # Title slide
    slide1.shapes.title.text = f"Insights for {client_name}"

    # Add Energy Summary Table
    slide2 = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide2.shapes.title.text = "Monthly kW, kWh Comparison & Load Factor​"
    add_table_to_slide(slide2, energy_summary, len(energy_summary) + 1,
                       len(energy_summary.columns), Inches(0.4), Inches(2.0), Inches(13), Inches(5))

    # Add other slides for plots
    slide3 = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide3.shapes.title.text = "Monthly Energy Consumption (kWh) and Peak Demand (kW)"
    slide3.shapes.add_picture(image_paths['energy_consumption'], Inches(0.4), Inches(2), Inches(12), Inches(5))

    slide4 = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide4.shapes.title.text = "Load Graph"
    slide4.shapes.add_picture(image_paths['energy_behavior'], Inches(0.4), Inches(2), Inches(13), Inches(5))

    slide5 = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide5.shapes.title.text = "Daily Average Demand"
    slide5.shapes.add_picture(image_paths['hourly_load_curve'], Inches(0.4), Inches(2), Inches(8), Inches(5.2))
    slide5.shapes.add_picture(image_paths['peak_demand_plot'], Inches(7), Inches(2), Inches(5), Inches(5.2))

    slide6 = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide6.shapes.title.text = "Hourly Averages Demand​"
    slide6.shapes.add_picture(image_paths['daily_kw_report'], Inches(0.4), Inches(2), Inches(8), Inches(5.2))
    slide6.shapes.add_picture(image_paths['demand_consumption_wesm'], Inches(7), Inches(2), Inches(5), Inches(5.2))

    # Save the PowerPoint
    output_path = os.path.join("generated_ppts", f"{client_name}_insights.pptx")
    ppt.save(output_path)

    return output_path

  except Exception as e:
    raise Exception(f"Error generating PowerPoint: {str(e)}")
